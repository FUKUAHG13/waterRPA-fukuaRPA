"""Excel 后端抽象与两种实现。

* ``OpenpyxlBackend``：基于 ``openpyxl`` 读写 .xlsx（惰性导入，未安装也能被
  import / 被 fukuaRPA 校验，只是运行动作时才提示安装）。
* ``MockExcelBackend``：纯内存工作簿，用于单元测试与无 openpyxl 环境的逻辑冒烟。

注意：openpyxl 只支持 .xlsx（不支持旧版 .xls）。若需 .xls，应另接 xlrd/转换，
不在本插件范围内。
"""

from __future__ import annotations

import contextlib
import os
import re
from abc import ABC, abstractmethod
from typing import TYPE_CHECKING, Any

if TYPE_CHECKING:  # 仅类型检查期可见：pdf_backend 在函数体内惰性导入以避免导入环
    from .pdf_backend import PdfBackend


def cell_to_rc(cell: str) -> tuple[int, int]:
    """把 "A1" 解析成 (1-based 行, 1-based 列)。"""

    m = re.match(r"^([A-Za-z]+)(\d+)$", cell.strip())
    if not m:
        raise ValueError(f"非法单元格地址：{cell!r}")
    col_letters, row = m.group(1).upper(), int(m.group(2))
    col = 0
    for ch in col_letters:
        col = col * 26 + (ord(ch) - ord("A") + 1)
    if row < 1:
        raise ValueError(f"非法单元格地址（行号必须 >= 1）: {cell}")
    return row, col


def _col_to_letter(col: int) -> str:
    """把 1-based 列号转成列字母：1->A, 26->Z, 27->AA（>26 列亦可）。"""

    if col <= 0:
        raise ValueError(f"非法列号：{col}")
    s = ""
    while col > 0:
        col, r = divmod(col - 1, 26)
        s = chr(ord("A") + r) + s
    return s


class ExcelBackend(ABC):
    # 读类动作是否需要文件真实存在；MockExcelBackend 覆盖为 False 以支持虚拟路径。
    requires_real_file: bool = True

    @abstractmethod
    def open(self, path: str) -> None:
        """加载工作簿（读或写前调用）。"""

    @abstractmethod
    def read_cell(self, sheet: str, cell: str) -> Any: ...

    @abstractmethod
    def read_range(self, sheet: str, start: str, end: str) -> list[list[Any]]: ...

    @abstractmethod
    def write_cell(self, sheet: str, cell: str, value: Any) -> None: ...

    @abstractmethod
    def write_range(self, sheet: str, start: str, values: list[list[Any]]) -> None: ...

    @abstractmethod
    def find(self, sheet: str, value: str, mode: str) -> list[str]:
        """返回命中单元格地址列表（如 ["A2","B2"]）。mode ∈ {contains,equals}。"""

    @abstractmethod
    def sheet_names(self) -> list[str]: ...

    @abstractmethod
    def row_count(self, sheet: str) -> int: ...

    @abstractmethod
    def save(self, path: str | None = None) -> None: ...

    @abstractmethod
    def close(self) -> None: ...


class OpenpyxlBackend(ExcelBackend):
    def __init__(self) -> None:
        self._wb = None  # 含公式的工作簿：用于写入与「计算值缺失」时的读取回退
        self._wb_data = None  # data_only 工作簿：用于读取缓存的计算值
        self._path: str | None = None

    def open(self, path: str) -> None:
        from openpyxl import Workbook, load_workbook  # 惰性

        # 释放上一次打开的工作簿句柄，避免操作间句柄泄漏（P2-1）。
        if self._wb is not None or self._wb_data is not None:
            self.close()

        self._path = path
        if os.path.exists(path):
            # 含公式版本：写入与读取回退使用（保留公式）。
            self._wb = load_workbook(path)
            # data_only 版本（读取 Excel 缓存的计算结果）改为「惰性加载」：
            # 纯写动作不再无谓地双加载整本工作簿（修复 P3 性能/资源缺陷）。
            self._wb_data = None
        else:
            # 写模式面向尚不存在的文件：新建工作簿，后续 save() 落地。
            self._wb = Workbook()
            self._wb_data = None

    def _ensure_data(self):
        """惰性加载 data_only 工作簿（仅在首次读取计算值时才真正打开）。"""
        if self._wb_data is None and self._path and os.path.exists(self._path):
            from openpyxl import load_workbook  # 惰性

            self._wb_data = load_workbook(self._path, data_only=True)
        return self._wb_data

    def _ensure(self):
        if self._wb is None:
            raise RuntimeError("尚未打开工作簿：请先 open(path)")

    def _sheet(self, name: str):
        self._ensure()
        if name:
            return self._wb[name]
        return self._wb.active

    def _data_sheet(self, name: str):
        ds = self._ensure_data()
        if ds is None:
            return None
        if name:
            return ds[name]
        return ds.active

    def read_cell(self, sheet: str, cell: str) -> Any:
        # 优先返回公式的计算值；openpyxl 限制——从未被 Excel 计算过的公式缓存值为
        # None，此时回退到含公式版本的值，避免静默返回 None。
        ds = self._data_sheet(sheet)
        val = ds[cell].value if ds is not None else None
        if val is None:
            ws = self._sheet(sheet)
            val = ws[cell].value
        return val

    def read_range(self, sheet: str, start: str, end: str) -> list[list[Any]]:
        ws = self._sheet(sheet)
        ds = self._data_sheet(sheet)
        out: list[list[Any]] = []
        for row in ws[start:end]:
            row_vals: list[Any] = []
            for c in row:
                val = ds[c.coordinate].value if ds is not None else None
                if val is None:
                    val = c.value
                row_vals.append(val)
            out.append(row_vals)
        return out

    def write_cell(self, sheet: str, cell: str, value: Any) -> None:
        ws = self._sheet(sheet)
        ws[cell] = value

    def write_range(self, sheet: str, start: str, values: list[list[Any]]) -> None:
        ws = self._sheet(sheet)
        row0, col0 = cell_to_rc(start)
        for r, row_vals in enumerate(values):
            for c, val in enumerate(row_vals):
                ws.cell(row=row0 + r, column=col0 + c, value=val)

    def find(self, sheet: str, value: str, mode: str) -> list[str]:
        ws = self._sheet(sheet)
        # 优先从 data_only 工作簿读取计算值进行匹配，避免命中公式文本（修复 D-11）。
        ds = self._data_sheet(sheet)
        target = str(value)
        hits: list[str] = []
        for row in ws.iter_rows():
            for cell in row:
                val = ds[cell.coordinate].value if ds is not None else None
                if val is None:
                    val = cell.value
                if val is None:
                    continue
                text = str(val)
                if mode == "equals" and text == target or mode != "equals" and target in text:
                    hits.append(cell.coordinate)
        return hits

    def sheet_names(self) -> list[str]:
        self._ensure()
        return list(self._wb.sheetnames)

    def row_count(self, sheet: str) -> int:
        ws = self._sheet(sheet)
        return ws.max_row or 0

    def save(self, path: str | None = None) -> None:
        self._ensure()
        target = path or self._path
        if not target:
            raise RuntimeError("save 需要目标路径（新建文件请显式传入 path）")
        # 原子写（修复 P1-③）：先写入同目录临时文件，再 os.replace 原子替换目标，
        # 避免写出阶段异常（被杀/断电/磁盘满）把原文件截断为 0 字节损坏。
        # os.replace 在同卷上是原子的，原文件在写入完成前不被触碰。
        import tempfile

        target_dir = os.path.dirname(os.path.abspath(target))
        os.makedirs(target_dir, exist_ok=True)
        fd, tmp = tempfile.mkstemp(dir=target_dir, prefix=".office_docs_", suffix=".tmp")
        os.close(fd)
        try:
            self._wb.save(tmp)
            os.replace(tmp, target)
        except BaseException:
            # 写出或替换失败：清理临时文件，绝不动原文件（修复 D-11 临时文件泄漏）。
            try:
                if os.path.exists(tmp):
                    os.remove(tmp)
            except OSError:
                pass
            raise
        # 保存后磁盘上的 data_only 副本可能已过期，失效以便后续读取重新惰性加载。
        self._wb_data = None

    def close(self) -> None:
        # 真正释放工作簿与其底层 ZipFile（修复 P3 句柄未关闭缺陷）。
        if self._wb is not None:
            with contextlib.suppress(Exception):
                self._wb.close()
            self._wb = None
        if self._wb_data is not None:
            with contextlib.suppress(Exception):
                self._wb_data.close()
            self._wb_data = None


class MockExcelBackend(ExcelBackend):
    """纯内存工作簿，预置一份示例数据便于测试。"""

    requires_real_file = False  # Mock 使用虚拟路径，无需磁盘文件存在

    def __init__(self) -> None:
        self._book: dict[str, list[list[Any]]] = {
            "Sheet1": [
                ["姓名", "年龄", "城市"],
                ["Alice", 30, "北京"],
                ["Bob", 25, "上海"],
                ["Carol", 28, "广州"],
            ]
        }
        self._opened: str | None = None

    def open(self, path: str) -> None:
        self._opened = path  # Mock 不真正读盘

    def _grid(self, sheet: str) -> list[list[Any]]:
        if sheet and sheet in self._book:
            return self._book[sheet]
        if sheet:
            self._book[sheet] = []
        # 默认取首个表
        if not self._book:
            self._book["Sheet1"] = []
        name = sheet or next(iter(self._book))
        return self._book[name]

    def read_cell(self, sheet: str, cell: str) -> Any:
        grid = self._grid(sheet)
        r, c = cell_to_rc(cell)
        if r - 1 >= len(grid) or c - 1 >= len(grid[r - 1]):
            return None
        return grid[r - 1][c - 1]

    def read_range(self, sheet: str, start: str, end: str) -> list[list[Any]]:
        grid = self._grid(sheet)
        r1, c1 = cell_to_rc(start)
        r2, c2 = cell_to_rc(end)
        out: list[list[Any]] = []
        for r in range(r1, r2 + 1):
            row_vals: list[Any] = []
            for c in range(c1, c2 + 1):
                if r - 1 < len(grid) and c - 1 < len(grid[r - 1]):
                    row_vals.append(grid[r - 1][c - 1])
                else:
                    row_vals.append(None)
            out.append(row_vals)
        return out

    def write_cell(self, sheet: str, cell: str, value: Any) -> None:
        grid = self._grid(sheet)
        r, c = cell_to_rc(cell)
        while len(grid) < r:
            grid.append([])
        while len(grid[r - 1]) < c:
            grid[r - 1].append(None)
        grid[r - 1][c - 1] = value

    def write_range(self, sheet: str, start: str, values: list[list[Any]]) -> None:
        grid = self._grid(sheet)
        r0, c0 = cell_to_rc(start)
        for r, row_vals in enumerate(values):
            rr = r0 + r
            while len(grid) < rr:
                grid.append([])
            for c, val in enumerate(row_vals):
                cc = c0 + c
                while len(grid[rr - 1]) < cc:
                    grid[rr - 1].append(None)
                grid[rr - 1][cc - 1] = val

    def find(self, sheet: str, value: str, mode: str) -> list[str]:
        grid = self._grid(sheet)
        target = str(value)
        hits: list[str] = []
        for r, row in enumerate(grid, start=1):
            for c, val in enumerate(row, start=1):
                if val is None:
                    continue
                text = str(val)
                if mode == "equals" and text == target or mode != "equals" and target in text:
                    hits.append(f"{_col_to_letter(c)}{r}")
        return hits

    def sheet_names(self) -> list[str]:
        return list(self._book.keys())

    def row_count(self, sheet: str) -> int:
        return len(self._grid(sheet))

    def save(self, path: str | None = None) -> None:
        # Mock 不落盘；真实 OpenpyxlBackend 会写文件。
        return None

    def close(self) -> None:
        self._opened = None


def create_excel_backend(name: str = "openpyxl") -> ExcelBackend:
    if name == "mock":
        return MockExcelBackend()
    if name in ("openpyxl", "xlsx"):
        return OpenpyxlBackend()
    raise ValueError(f"未知 Excel 后端：{name!r}（可选：openpyxl / mock）")


# --------------------------------------------------------------------------
# 进程级单例缓存与生命周期管理
# --------------------------------------------------------------------------

# 进程级单例缓存：openpyxl/pypdf 首次构造有开销，避免每个动作重开。
_XL_CACHE: dict[str, ExcelBackend] = {}
_PDF_CACHE: dict[str, PdfBackend] = {}


def get_excel_backend(name: str = "openpyxl", **kwargs) -> ExcelBackend:
    """按名字取 Excel 后端，进程内复用同一实例（单例）。"""
    cached = _XL_CACHE.get(name)
    if cached is None:
        cached = create_excel_backend(name, **kwargs)
        _XL_CACHE[name] = cached
    return cached


def get_pdf_backend(name: str = "pypdf", **kwargs) -> PdfBackend:
    """按名字取 PDF 后端，进程内复用同一实例（单例）。"""
    cached = _PDF_CACHE.get(name)
    if cached is None:
        from .pdf_backend import create_pdf_backend  # 惰性，避免任何导入环
        cached = create_pdf_backend(name, **kwargs)
        _PDF_CACHE[name] = cached
    return cached


def close_all() -> None:
    """释放所有缓存的 Excel/PDF 后端并清空缓存。

    生命周期：本模块持有进程级单例缓存；宿主（P3）退出时通过 ``atexit`` 自动调用
    本函数，关闭工作簿句柄/释放资源。幂等——重复调用安全。
    """
    for backend in list(_XL_CACHE.values()) + list(_PDF_CACHE.values()):
        with contextlib.suppress(Exception):
            backend.close()
    _XL_CACHE.clear()
    _PDF_CACHE.clear()


try:
    import atexit
    atexit.register(close_all)
except Exception:  # noqa: BLE001 极少数无 atexit 的环境（如受限嵌入式解释器）
    pass
